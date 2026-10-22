"""Turns detailedfoil's stock export into the Product Search list.

    python3 .github/product-search/make-list.py path/to/export.xlsx
    python3 .github/product-search/make-list.py path/to/export.txt --add

The export can be the workbook (the "Paste Here!" sheet, table tblSKU) or the report copied as text (tab separated,
each heading on a line of its own, then one product per line).
--add keeps every product already in the list exactly as it is and only adds codes it hasn't got (detailedfoil,
22 Oct 2026: "just add the extra items, dont need to delete anything"). Without --add the list is replaced.

Writes tools/product-search/data/products.json: one product per line, [SKU, name, flags], where flags adds up
1 = hazardous, 2 = aerosol, 4 = non-trolleyable, 8 = sharp or liquid (a column the export hasn't got counts as no).
Stock, locations and anything else store-specific are left out on purpose: the repo is public.
The workbook needs openpyxl (pip install openpyxl). It reads the values Excel saved, so save the workbook in Excel first.
"""
import json, re, sys, datetime, pathlib

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT = ROOT / 'tools/product-search/data/products.json'
SKU = re.compile(r'^(\d{3}[A-Z]{2}|\d{4}[A-Z]|\d{5})$')
FLAGS = [('Hazardous', 1), ('Aerosol', 2), ('Non Trolleyable', 4), ('Sharp or Liquid', 8)]


def code(v):
    return str(v or '').strip().upper().replace('O', '0')


def rows_xlsx(src):
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb['Paste Here!'] if 'Paste Here!' in wb.sheetnames else wb.worksheets[0]
    head = None
    for row in ws.iter_rows(values_only=True):
        if head is None:
            if row and 'SKU Code' in row: head = [str(c or '').strip() for c in row]
            continue
        yield head, row


def rows_text(src):
    # headings one per line (the first with its tooltips after a tab), then tab separated products
    lines = pathlib.Path(src).read_text(encoding='utf-8', errors='replace').splitlines()
    head, i = None, 0
    while i < len(lines) and lines[i].split('\t')[0].strip() != 'SKU Code': i += 1
    if i == len(lines): return
    head = []
    for line in lines[i:]:
        cells = line.split('\t')
        if head and len(cells) > 1 and SKU.match(code(cells[0])): break
        head.append(cells[0].strip())
        i += 1
    for line in lines[i:]:
        if line.strip(): yield head, line.split('\t')


def main(src, add):
    read = rows_xlsx if src.lower().endswith(('.xlsx', '.xlsm')) else rows_text
    items, seen, bad, cols, missing = [], set(), [], None, []
    for head, row in read(src):
        if cols is None:
            cols = {name: head.index(name) for name in head if name}
            if 'SKU Code' not in cols or 'Description' not in cols: sys.exit('Missing columns: SKU Code and Description are needed')
            missing = [f for f, _ in FLAGS if f not in cols]
        cell = lambda name: row[cols[name]] if name in cols and cols[name] < len(row) else ''
        sku = code(cell('SKU Code'))
        if not sku: continue
        if not SKU.match(sku): bad.append(sku); continue
        if sku in seen: continue
        seen.add(sku)
        name = re.sub(r'\s+', ' ', str(cell('Description') or '')).strip()
        flags = sum(bit for col, bit in FLAGS if str(cell(col) or '').strip().upper() == 'YES')
        items.append([sku, name, flags])
    if cols is None: sys.exit('No "SKU Code" heading found')
    note = []
    if add:
        old = json.loads(OUT.read_text(encoding='utf-8'))['items']
        have = {i[0] for i in old}
        new = [i for i in items if i[0] not in have]
        note.append(f'{len(new)} new added to the {len(old)} already there'
                    f' ({sum(1 for c in have if c not in seen)} not in this export, kept)')
        items = old + new
    items.sort(key=lambda x: x[0])
    version = datetime.date.today().isoformat()
    lines = ',\n'.join('    ' + json.dumps(i, ensure_ascii=False) for i in items)
    OUT.write_text('{\n  "version": %s,\n  "count": %d,\n  "flags": {"1": "Hazardous", "2": "Aerosol", "4": "Non-trolleyable", "8": "Sharp or liquid"},\n  "items": [\n%s\n  ]\n}\n'
                   % (json.dumps(version), len(items), lines), encoding='utf-8')
    if missing: note.append('no ' + ', '.join(missing) + ' column, counted as no')
    if bad: note.append(f'skipped {len(bad)} odd codes: {bad[:10]}')
    print(f'{len(items)} products written to {OUT.relative_to(ROOT)}' + ''.join('; ' + n for n in note))


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if a != '--add']
    if len(args) != 1: sys.exit(__doc__)
    main(args[0], '--add' in sys.argv[1:])
